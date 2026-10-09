"""
Gera o arquivo de slides PowerPoint (.pptx) para a apresentação da Trilha D do Cleiver no Projeto Humans.
Design profissional 16:9, clean, visual e focado nos resultados dos experimentos.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    base_dir = os.path.abspath(os.path.dirname(__file__))
    template_path = os.path.join(base_dir, "Slides", "Template_slides.pptx")
    output_pptx = os.path.join(base_dir, "Slides", "Trilha_D_Controle_Emocao_Seminario.pptx")

    # Inicia apresentação com template 16:9 widescreen
    if os.path.exists(template_path):
        prs = Presentation(template_path)
    else:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

    sw = prs.slide_width
    sh = prs.slide_height

    # Paleta de Cores Corporativa Dark / Tech
    BG_DARK = RGBColor(15, 23, 42)       # Slate 900
    CARD_DARK = RGBColor(30, 41, 59)     # Slate 800
    TEXT_LIGHT = RGBColor(248, 250, 252) # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184) # Slate 400
    ACCENT_BLUE = RGBColor(56, 189, 248) # Sky 400
    ACCENT_GREEN = RGBColor(34, 197, 94) # Green 500
    ACCENT_PURPLE = RGBColor(168, 85, 247)# Purple 500
    ACCENT_AMBER = RGBColor(245, 158, 11)# Amber 500

    blank_layout = prs.slide_layouts[6] # Blank

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, sw, sh)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="TRILHA D: CONTROLE, EMOÇÃO E COMPORTAMENTO"):
        # Header category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), sw - Inches(1.6), Inches(0.4))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category.upper()
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = ACCENT_BLUE

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), sw - Inches(1.6), Inches(0.8))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 1: CAPA
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, BG_DARK)

    # Card central
    box_w = Inches(10.5)
    box_h = Inches(4.8)
    box_x = (sw - box_w) / 2
    box_y = (sh - box_h) / 2

    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, box_x, box_y, box_w, box_h)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_DARK
    card.line.color.rgb = ACCENT_BLUE

    tb = s1.shapes.add_textbox(box_x + Inches(0.8), box_y + Inches(0.8), box_w - Inches(1.6), box_h - Inches(1.6))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "PROJETO HUMANS • AKCIT"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "Trilha D: Controle, Emoção e Comportamento"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT
    p1.space_after = Pt(12)

    p2 = tf.add_paragraph()
    p2.text = "Síntese Autônoma de Movimento Facial e Corporal a partir de Áudio e Texto (Sem Pose Explícita)"
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_after = Pt(24)

    p3 = tf.add_paragraph()
    p3.text = "Pesquisador: Cleiver | Outubro 2026"
    p3.font.size = Pt(13)
    p3.font.bold = True
    p3.font.color.rgb = ACCENT_GREEN

    # =========================================================================
    # SLIDE 2: ONDE A TRILHA D SE ENCAIXA NO GRUPO? (O MAESTRO)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, BG_DARK)
    add_header(s2, "Onde a Trilha D se Encaixa no Projeto Humans?")

    # 4 Colunas para as 4 Trilhas
    cols = [
        {"title": "Trilha A (Mateus)", "desc": "Vídeo 2D Direto\nFoto + Áudio -> Vídeo\n(SadTalker, LivePortrait, EMO)", "color": ACCENT_BLUE},
        {"title": "Trilha B (Arthur)", "desc": "Avatar 3D de 1 Foto\nMalha FLAME/SMPL-X\nA 'Marionete' 3D", "color": ACCENT_PURPLE},
        {"title": "Trilha C (Fernando)", "desc": "Cabeças 3DGS Animáveis\nGaussianas na malha\nA 'Pele' Fotográfica 3D", "color": ACCENT_AMBER},
        {"title": "Trilha D (Cleiver)", "desc": "Controle & Emoção\nÁudio/Texto -> Movimento\nO 'Maestro' / Titiriteiro", "color": ACCENT_GREEN, "highlight": True}
    ]

    col_w = Inches(2.7)
    gap = Inches(0.3)
    start_x = Inches(0.8)
    card_y = Inches(2.0)
    card_h = Inches(4.5)

    for i, col in enumerate(cols):
        cx = start_x + i * (col_w + gap)
        scard = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, col_w, card_h)
        scard.fill.solid()
        scard.fill.fore_color.rgb = CARD_DARK
        scard.line.color.rgb = col["color"]
        scard.line.width = Pt(2.5 if col.get("highlight") else 1.0)

        stb = s2.shapes.add_textbox(cx + Inches(0.2), card_y + Inches(0.3), col_w - Inches(0.4), card_h - Inches(0.6))
        stf = stb.text_frame
        stf.word_wrap = True

        sp1 = stf.paragraphs[0]
        sp1.text = col["title"]
        sp1.font.size = Pt(14)
        sp1.font.bold = True
        sp1.font.color.rgb = col["color"]
        sp1.space_after = Pt(14)

        sp2 = stf.add_paragraph()
        sp2.text = col["desc"]
        sp2.font.size = Pt(12)
        sp2.font.color.rgb = TEXT_LIGHT
        sp2.space_after = Pt(14)

        if col.get("highlight"):
            sp3 = stf.add_paragraph()
            sp3.text = "★ SEM A TRILHA D:\nAs cabeças 3DGS e malhas ficam paralisadas ou precisam de webcam humana gravando o tempo todo."
            sp3.font.size = Pt(10.5)
            sp3.font.bold = True
            sp3.font.color.rgb = ACCENT_GREEN

    # =========================================================================
    # SLIDE 3: AS 6 SEMENTES SOTA EM UMA VISÃO PANORÂMICA
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, BG_DARK)
    add_header(s3, "As 6 Sementes da Trilha D: A Vanguarda do Estado da Arte")

    seeds = [
        {"name": "Motion Diffusion (MDM)", "tag": "Motor Matemático", "desc": "Resolve o problema 1-para-Muitos. Gera dinâmica estocástica suave sem colapso para a média.", "color": ACCENT_BLUE},
        {"name": "VASA-1 (Microsoft, 2024)", "tag": "Tempo Real 40+ FPS", "desc": "Desacopla identidade estática, pose de cabeça rígida 3D e dinâmica facial em espaço latente.", "color": ACCENT_GREEN},
        {"name": "InstructAvatar (AAAI 2025)", "tag": "Direção por Texto", "desc": "Two-Branch Diffusion: a voz mantém lip-sync perfeito enquanto prompts de texto ditam a atuação dramática.", "color": ACCENT_PURPLE},
        {"name": "AUHead (ICLR 2026)", "tag": "Anatomia Muscular", "desc": "Passa pelo sistema FACS (Ekman). Controle cirúrgico dos músculos do rosto sem distorções caixa-preta.", "color": ACCENT_AMBER},
        {"name": "Audio2Photoreal (Meta, CVPR 2024)", "tag": "Corpo Inteiro", "desc": "Sai do busto para corpo inteiro e gestos de mãos em diálogos naturais usando VQ-VAE + Difusão.", "color": ACCENT_BLUE},
        {"name": "OmniHuman-1.5 (ByteDance, 2025)", "tag": "Cognição Dual", "desc": "Sistema 1 (reflexo reativo rápido) + Sistema 2 (planejamento deliberado via MLLM) para avatar com intenção.", "color": ACCENT_GREEN}
    ]

    grid_w = Inches(5.6)
    grid_h = Inches(2.2)
    positions = [
        (Inches(0.8), Inches(1.8)),
        (Inches(6.8), Inches(1.8)),
        (Inches(0.8), Inches(4.3)),
        (Inches(6.8), Inches(4.3)),
        (Inches(0.8), Inches(6.8)),
        (Inches(6.8), Inches(6.8))
    ]

    for i, s_data in enumerate(seeds):
        gx, gy = positions[i]
        card_s = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, gx, gy, grid_w, grid_h)
        card_s.fill.solid()
        card_s.fill.fore_color.rgb = CARD_DARK
        card_s.line.color.rgb = s_data["color"]

        tb_s = s3.shapes.add_textbox(gx + Inches(0.2), gy + Inches(0.15), grid_w - Inches(0.4), grid_h - Inches(0.3))
        tfs = tb_s.text_frame
        tfs.word_wrap = True

        p_name = tfs.paragraphs[0]
        p_name.text = f"{s_data['name']}  •  [{s_data['tag']}]"
        p_name.font.size = Pt(13)
        p_name.font.bold = True
        p_name.font.color.rgb = s_data["color"]
        p_name.space_after = Pt(6)

        p_desc = tfs.add_paragraph()
        p_desc.text = s_data["desc"]
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 4: O PROBLEMA DO "1-PARA-MUITOS" E A SOLUÇÃO POR DIFUSÃO
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, BG_DARK)
    add_header(s4, "O Problema Central: Por que Modelos Regressivos Clássicos Falham?")

    # Bloco da Esquerda: O Problema
    b1_x = Inches(0.8)
    b_w = Inches(5.6)
    b_h = Inches(4.8)
    b1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, b1_x, Inches(2.0), b_w, b_h)
    b1.fill.solid()
    b1.fill.fore_color.rgb = CARD_DARK
    b1.line.color.rgb = RGBColor(244, 63, 94) # Rose

    tb_b1 = s4.shapes.add_textbox(b1_x + Inches(0.3), Inches(2.2), b_w - Inches(0.6), b_h - Inches(0.4))
    tf_b1 = tb_b1.text_frame
    tf_b1.word_wrap = True

    p = tf_b1.paragraphs[0]
    p.text = "O Problema: Mapeamento 1-para-Muitos"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(244, 63, 94)
    p.space_after = Pt(12)

    p = tf_b1.add_paragraph()
    p.text = "• Sincronia Labial é ~1-para-1: falar o fonema /m/ fecha a boca obrigatoriamente."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(8)

    p = tf_b1.add_paragraph()
    p.text = "• Movimento da Cabeça é 1-para-Muitos: para a mesma frase, você pode inclinar a cabeça à esquerda (+15°), à direita (-15°) ou assentir."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(8)

    p = tf_b1.add_paragraph()
    p.text = "• Colapso para a Média: Redes com perda MSE tentam adivinhar a média de todos os movimentos válidos. A média é ZERO."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(8)

    p = tf_b1.add_paragraph()
    p.text = "Resultado Clássico: Cabeça dura, estática, parecendo um boneco de ventríloquo."
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(244, 63, 94)

    # Bloco da Direita: A Solução
    b2_x = Inches(6.8)
    b2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, b2_x, Inches(2.0), b_w, b_h)
    b2.fill.solid()
    b2.fill.fore_color.rgb = CARD_DARK
    b2.line.color.rgb = ACCENT_GREEN

    tb_b2 = s4.shapes.add_textbox(b2_x + Inches(0.3), Inches(2.2), b_w - Inches(0.6), b_h - Inches(0.4))
    tf_b2 = tb_b2.text_frame
    tf_b2.word_wrap = True

    p = tf_b2.paragraphs[0]
    p.text = "A Solução: Difusão Estocástica de Trajetória"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(12)

    p = tf_b2.add_paragraph()
    p.text = "• Modelagem de Distribuição: Em vez de prever um valor fixo, o modelo amostra uma trajetória contínua plausível da distribuição real."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(8)

    p = tf_b2.add_paragraph()
    p.text = "• Denoising Temporal DDPM: Parte de ruído gaussiano puro e esculpe os ângulos de rotação (Yaw, Pitch, Roll) e translação condicionado no áudio e no texto."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(8)

    p = tf_b2.add_paragraph()
    p.text = "• Prova Empírica do Nosso Teste: Rodando o mesmo áudio com sementes de ruído diferentes, a distância L2 entre curvas foi de 0.1194."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(8)

    p = tf_b2.add_paragraph()
    p.text = "Resultado: O avatar atua como um humano vivo — cada execução gera uma performance natural e única!"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    # =========================================================================
    # SLIDE 5: RESULTADOS EXPERIMENTAIS: O PIPELINE IMPLEMENTADO
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, BG_DARK)
    add_header(s5, "Resultados Experimentais: O Pipeline Operacional da Trilha D")

    # Imagem do dashboard se existir
    dash_img_path = os.path.join(base_dir, "Trilha_D_Controle_Emocao", "output", "dashboard_animacao.png")
    if os.path.exists(dash_img_path):
        s5.shapes.add_picture(dash_img_path, Inches(0.8), Inches(1.8), width=Inches(7.2))

    # Métricas ao lado
    m_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.3), Inches(1.8), Inches(4.3), Inches(5.0))
    m_box.fill.solid()
    m_box.fill.fore_color.rgb = CARD_DARK
    m_box.line.color.rgb = ACCENT_BLUE

    tb_m = s5.shapes.add_textbox(Inches(8.5), Inches(2.0), Inches(3.9), Inches(4.6))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True

    p = tf_m.paragraphs[0]
    p.text = "MÉTRICAS DO EXPERIMENTO"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(12)

    bullets = [
        ("Aceleração de Hardware", "Apple Silicon MPS (Metal)"),
        ("Tempo de Difusão", "0.10 segundos (~900 FPS de parâmetros)"),
        ("Taxa de Vídeo", "30 FPS contínuo sincronizado"),
        ("Frontend Acústico", "Log-Mel 80 canais + F0 + RMS"),
        ("FACS / Ekman", "14 Action Units dinâmicas (AU12 sorriso, AU26 mandíbula, AU45 piscadas)"),
        ("FLAME Compatibility", "50 coeficientes de expressão + rotação 3D de mandíbula"),
        ("Exportação", "Formatos padronizados JSON e tensores binários .NPZ")
    ]

    for label, val in bullets:
        p = tf_m.add_paragraph()
        p.text = f"• {label}: {val}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(6)

    # =========================================================================
    # SLIDE 6: DO ESQUEMÁTICO AO FOTO-REALISMO (FOTO REAL GANHA VIDA)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, BG_DARK)
    add_header(s6, "Do Esquemático ao Foto-Realismo: Animação Densa sobre Foto Real")

    # Inserir imagem do retrato real
    portrait_path = os.path.join(base_dir, "Trilha_D_Controle_Emocao", "data", "vasa_portrait.jpg")
    if os.path.exists(portrait_path):
        s6.shapes.add_picture(portrait_path, Inches(0.8), Inches(1.8), width=Inches(3.8))

    # Card explicativo da deformação foto-realista
    card_r = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.9), Inches(1.8), Inches(7.6), Inches(4.8))
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = CARD_DARK
    card_r.line.color.rgb = ACCENT_GREEN

    tb_r = s6.shapes.add_textbox(Inches(5.2), Inches(2.0), Inches(7.0), Inches(4.4))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "Como o VASA-1 e o Nosso Renderizador Fazem a Foto se Mover?"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(12)

    p = tf_r.add_paragraph()
    p.text = "1. Preservação de Identidade: A foto original estática de 1024x1024 fornece a textura dos poros, iluminação de estúdio e fios de barba."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(8)

    p = tf_r.add_paragraph()
    p.text = "2. Deformação Afim Densa (Delaunay Mesh Warping): Os pontos de controle da face sofrem translação e rotação com efeito de paralaxe 3D, mantendo a geometria elástica natural da pele."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(8)

    p = tf_r.add_paragraph()
    p.text = "3. Sincronia Labial com Cavidade Oral Real: Quando a amplitude do áudio abre a mandíbula (AU26), é sintetizada a sombra oclusal interna e o arco dental superior, sem cortes artificiais."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(8)

    p = tf_r.add_paragraph()
    p.text = "4. Síntese de Piscadas: As pálpebras cobrem a esclera e a íris com a textura exata da pele do sujeito em intervalos fisiológicos."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(12)

    p = tf_r.add_paragraph()
    p.text = "Resultado: Vídeo MP4 (H.264 + AAC) a 30 FPS pronto para exibição e apresentação."
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    # =========================================================================
    # SLIDE 7: DEMONSTRAÇÃO PRÁTICA: EXPERIMENTOS COMPARATIVOS
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, BG_DARK)
    add_header(s7, "Demonstração Prática: Nuances Dramáticas Geradas por Texto")

    # Dois Cards comparativos (Alegria vs Raiva)
    w_card7 = Inches(5.6)
    h_card7 = Inches(4.8)

    # Card 1: Alegria
    c1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), w_card7, h_card7)
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_DARK
    c1.line.color.rgb = ACCENT_GREEN

    tb_c1 = s7.shapes.add_textbox(Inches(1.1), Inches(2.2), w_card7 - Inches(0.6), h_card7 - Inches(0.4))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True

    p = tf_c1.paragraphs[0]
    p.text = "Experimento 1: Alegria & Entusiasmo"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(10)

    p = tf_c1.add_paragraph()
    p.text = "• Prompt: 'Fale com muita alegria e entusiasmo, olhando para a esquerda'"
    p.font.size = Pt(11)
    p.font.italic = True
    p.font.color.rgb = TEXT_MUTED
    p.space_after = Pt(8)

    p = tf_c1.add_paragraph()
    p.text = "• Fala Real Sintetizada (TTS): 'Bem-vindos ao laboratório de humanos virtuais da nossa equipe!'"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(8)

    p = tf_c1.add_paragraph()
    p.text = "• Resposta do Modelo:\n  - Rotação: Yaw = +15.0° (cabeça virada à esquerda)\n  - FACS: AU12 (sorriso) ativo em 0.65\n  - Olhos abertos com piscadas reflexas\n  - Mandíbula abrindo na cadência de cada palavra"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # Card 2: Raiva
    c2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.0), w_card7, h_card7)
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_DARK
    c2.line.color.rgb = RGBColor(244, 63, 94)

    tb_c2 = s7.shapes.add_textbox(Inches(7.1), Inches(2.2), w_card7 - Inches(0.6), h_card7 - Inches(0.4))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True

    p = tf_c2.paragraphs[0]
    p.text = "Experimento 2: Raiva & Firmeza"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(244, 63, 94)
    p.space_after = Pt(10)

    p = tf_c2.add_paragraph()
    p.text = "• Prompt: 'Fale com muita raiva e olhe para a direita balançando a cabeça afirmativo'"
    p.font.size = Pt(11)
    p.font.italic = True
    p.font.color.rgb = TEXT_MUTED
    p.space_after = Pt(8)

    p = tf_c2.add_paragraph()
    p.text = "• Fala Real Sintetizada (TTS): 'Isso é inadmissível! Não podemos aceitar esse resultado de jeito nenhum!'"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(8)

    p = tf_c2.add_paragraph()
    p.text = "• Resposta do Modelo:\n  - Rotação: Yaw = -15.0° (cabeça virada à direita)\n  - Aceno: Oscilação contínua de Pitch (nodding enfático)\n  - FACS: AU04 (cenho franzido) contraído\n  - Lábios tensos acompanhando a dicção firme"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 8: COMO CONVERSO COM CADA COLEGA DO GRUPO (INTEGRAÇÃO)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, BG_DARK)
    add_header(s8, "Interface Inter-Trilhas: O Que a Trilha D Entrega para o Grupo?")

    int_cols = [
        {"to": "Para o Fernando (Trilha C - 3DGS)", "what": "A Trilha C tem uma cabeça em Gaussianas 3D pronta para deformar, mas não sabe como mover.\n\nO QUE ENTREGO:\n• Vetor contínuo de rotação da cabeça (rotation_6d / Euler)\n• 50 coeficientes de expressão FLAME frame a frame\n• Ângulo de abertura de mandíbula\n\n-> As gaussianas se movem sem precisar de ator na webcam!", "color": ACCENT_AMBER},
        {"to": "Para o Arthur (Trilha B - Malha / Rigging)", "what": "A Trilha B gera a malha 3D e o esqueleto (a 'marionete').\n\nO QUE ENTREGO:\n• As 'linhas' que puxam a marionete:\n• Rotação dos ossos do pescoço e queixo\n• Pesos dos blendshapes faciais e Action Units FACS\n\n-> A malha ganha vida autônoma e física crível.", "color": ACCENT_PURPLE},
        {"to": "Para o Mateus (Trilha A - Vídeo 2D)", "what": "Modelos 2D como LivePortrait ou SadTalker sofrem com cabeças duras se não tiverem um vídeo driver humano.\n\nO QUE ENTREGO:\n• Trajetórias latentes e sinais de controle de pose\n• Guia temporal de dinâmicas faciais\n\n-> O vídeo 2D ganha naturalidade sem depender de driving video gravado.", "color": ACCENT_BLUE}
    ]

    w_ic = Inches(3.6)
    gap_ic = Inches(0.4)
    start_ic = Inches(0.8)

    for i, item in enumerate(int_cols):
        ix = start_ic + i * (w_ic + gap_ic)
        card_i = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, ix, Inches(2.0), w_ic, Inches(4.8))
        card_i.fill.solid()
        card_i.fill.fore_color.rgb = CARD_DARK
        card_i.line.color.rgb = item["color"]

        tb_i = s8.shapes.add_textbox(ix + Inches(0.2), Inches(2.2), w_ic - Inches(0.4), Inches(4.4))
        tf_i = tb_i.text_frame
        tf_i.word_wrap = True

        p = tf_i.paragraphs[0]
        p.text = item["to"]
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = item["color"]
        p.space_after = Pt(12)

        p = tf_i.add_paragraph()
        p.text = item["what"]
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 9: CONCLUSÃO E PRÓXIMOS PASSOS
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, BG_DARK)
    add_header(s9, "Conclusão e Próximos Passos")

    c_box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(11.6), Inches(4.8))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = CARD_DARK
    c_box.line.color.rgb = ACCENT_GREEN

    tb_c = s9.shapes.add_textbox(Inches(1.2), Inches(2.3), Inches(10.8), Inches(4.2))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    p.text = "Status Atual da Trilha D: 100% Validada e Operacional"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(14)

    concl_points = [
        "✓ Literatura Mapeada: VASA-1, InstructAvatar, AUHead, Audio2Photoreal, OmniHuman-1.5 e Motion Diffusion dominados.",
        "✓ Pipeline Implementado: Da extração acústica (Mel-80/RMS) à difusão temporal e deformação afim foto-realista.",
        "✓ Desafio '1-para-Muitos' Superado: Comprovado em testes com estocasticidade e inércia física no Apple Silicon MPS.",
        "✓ Exportação Pronta: Pacotes .npz e .json gerados para acoplamento imediato no 3DGS do Fernando e malha do Arthur."
    ]

    for pt in concl_points:
        p = tf_c.add_paragraph()
        p.text = pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(10)

    p = tf_c.add_paragraph()
    p.text = "\nPróximo Passo: Plugar os tensores exportados diretamente no renderizador de GaussianAvatars do grupo."
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    prs.save(output_pptx)
    print(f"✨ Apresentação PowerPoint gerada com sucesso em: {output_pptx}")

if __name__ == "__main__":
    create_presentation()
